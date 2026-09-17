#!/usr/bin/env python3
"""
Tests for the budget sync's expense side.

Plain assertions, no pytest — it is not installed and this repo has no test
suite to hang one off. Every test builds its own workbook in a temp directory,
so nothing here touches OneDrive, the real data workbook, or the network.

Run:
    python3 scripts/onboarding/test_budget_sync.py
"""

import sys
import tempfile
import traceback
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

import budget_sync as bs  # noqa: E402
from build_v2 import EXPENSE_COLS, INCOME_COLS  # noqa: E402

TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


def fixture_workbook(tmp, expenses=None, incomes=None):
    """
    A minimal data workbook with the sheets budget_sync reads.

    Vendors and categories are invented. This repository is public, so no
    fixture ever carries a real name.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, cols, rows in (
        ("Incomes", INCOME_COLS, incomes or []),
        ("Expenses", EXPENSE_COLS, expenses or []),
    ):
        ws = wb.create_sheet(sheet)
        for i, (name, _width) in enumerate(cols, start=1):
            ws.cell(1, i, name)
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
    path = tmp / "FRC Data test.xlsx"
    wb.save(path)
    return path


@test
def test_dump_sheet_reports_header_and_rows():
    import io
    from contextlib import redirect_stdout

    with tempfile.TemporaryDirectory() as td:
        path = fixture_workbook(Path(td), expenses=[
            [date(2026, 8, 1), "Acme Parts", "Robot parts", 120.5, "team card", "INV-1", "", ""],
        ])
        buf = io.StringIO()
        with redirect_stdout(buf):
            bs.dump_sheet(path, "Expenses")
        out = buf.getvalue()
        assert "Expenses" in out, out
        assert "'Vendor'" in out, "header cells should be printed with repr"
        assert "Acme Parts" in out, "first data row should be printed"
        assert "8 columns" in out, f"expected the column count, got:\n{out}"


@test
def test_read_expenses_roundtrips_values_and_order():
    rows_in = [
        [date(2026, 8, 1), "Acme Parts", "Robot parts", 120.5, "team card", "INV-1", "", ""],
        [date(2026, 8, 9), "Bolt Supply", "Raw materials", 40.0, "team card", "INV-2", "", "rush"],
    ]
    with tempfile.TemporaryDirectory() as td:
        path = fixture_workbook(Path(td), expenses=rows_in)
        wb = openpyxl.load_workbook(path, data_only=True)
        got = bs.read_expenses(wb)
    assert len(got) == 2, got
    assert [r[1] for r in got] == ["Acme Parts", "Bolt Supply"], "order must be preserved"
    assert got[0][3] == 120.5
    # The padding branch itself rarely fires on real openpyxl input — the
    # sheet's bounding box already supplies width 8 here — but the contract
    # callers rely on is the output width, not how it got there.
    assert all(len(r) == 8 for r in got), "every returned row is exactly the spec width"


@test
def test_read_expenses_skips_blank_rows_and_missing_sheet():
    with tempfile.TemporaryDirectory() as td:
        path = fixture_workbook(Path(td), expenses=[
            [date(2026, 8, 1), "Acme Parts", "Robot parts", 10.0, "", "", "", ""],
            [None, None, None, None, None, None, None, None],
        ])
        wb = openpyxl.load_workbook(path, data_only=True)
        assert len(bs.read_expenses(wb)) == 1
        del wb["Expenses"]
        assert bs.read_expenses(wb) == [], "a missing sheet is empty, not an error"


@test
def test_expense_ids_present_ignores_hand_entered_rows():
    rows = [
        [date(2026, 8, 1), "Acme Parts", "Robot parts", 10.0, "", "", "EXP-1", ""],
        [date(2026, 8, 2), "Bolt Supply", "Raw materials", 20.0, "", "", "", ""],
        [date(2026, 8, 3), "Cog Co", "Tools", 30.0, "", "", "  EXP-2  ", ""],
    ]
    assert bs.expense_ids_present(rows) == {"EXP-1", "EXP-2"}, bs.expense_ids_present(rows)


@test
def test_overwide_rows_are_detected():
    # openpyxl pads every row it returns to the sheet's widest row, so the
    # only reliable way to find the one overwide row is to read the actual
    # worksheet — and a blank row ahead of the offender proves the result is
    # a real (1-based) sheet row number, not an index into some filtered list.
    normal = [date(2026, 8, 1), "Acme Parts", "Robot parts", 10.0, "", "", "", ""]
    blank = [None] * 8
    offender = [date(2026, 8, 3), "Cog Co", "Tools", 30.0, "", "", "", "", "stray value"]
    with tempfile.TemporaryDirectory() as td:
        path = fixture_workbook(Path(td), expenses=[normal, blank, offender, normal])
        wb = openpyxl.load_workbook(path, data_only=True)
        got = bs.overwide_expense_rows(wb["Expenses"])
    assert got == [4], f"expected only the offending sheet row, got {got}"

    with tempfile.TemporaryDirectory() as td:
        path = fixture_workbook(Path(td), expenses=[normal, normal])
        wb = openpyxl.load_workbook(path, data_only=True)
        assert bs.overwide_expense_rows(wb["Expenses"]) == []


@test
def test_norm_category_collapses_case_and_spacing_only():
    from build_v2 import norm_category
    assert norm_category("Robot parts") == norm_category("robot parts")
    assert norm_category("Robot  Parts") == norm_category("Robot Parts")
    assert norm_category("  Tools ") == "tools"
    assert norm_category(None) == ""
    # norm() splits on the first hyphen, which would truncate this to 'pre'.
    # A category is not a name and must not be treated like one.
    assert norm_category("Pre-season build") == "pre-season build"
    assert norm_category("Travel") != norm_category("Tools")


@test
def test_category_counts_tallies_non_blank_categories():
    rows = [
        [date(2026, 8, 1), "Acme Parts", "Robot parts", 10.0, "", "", "", ""],
        [date(2026, 8, 2), "Bolt Supply", "Robot parts", 20.0, "", "", "", ""],
        [date(2026, 8, 3), "Cog Co", "Tools", 30.0, "", "", "", ""],
        [date(2026, 8, 4), "Dial Ltd", "", 40.0, "", "", "", ""],
        [date(2026, 8, 5), "Edge Inc", "   ", 50.0, "", "", "", ""],
    ]
    counts = bs.category_counts(rows)
    assert dict(counts) == {"Robot parts": 2, "Tools": 1}, dict(counts)


@test
def test_category_variants_groups_spellings_and_ignores_clean_ones():
    from collections import Counter
    counts = Counter({"Robot parts": 18, "robot parts": 2, "Robot Parts": 1, "Travel": 7})
    groups = bs.category_variants(counts)
    assert len(groups) == 1, groups
    assert [name for name, _n in groups[0]] == ["Robot parts", "robot parts", "Robot Parts"], groups[0]
    assert groups[0][0] == ("Robot parts", 18), "most-used spelling leads"


def cell_equal(before, after):
    """
    Compare a written cell with the one read back.

    openpyxl writes '' as an empty cell and reads it back as None, so a strict
    == would fail on every blank a person left blank. Both spellings mean
    empty. 0 is not empty and is compared normally.
    """
    if before in (None, "") and after in (None, ""):
        return True
    # A date-formatted cell always reads back as `datetime`, even one written
    # as a plain `date` — verified against actual openpyxl output, not
    # assumed. That is a serialization detail, not a changed value.
    if isinstance(before, date) and isinstance(after, datetime):
        return before == after.date()
    return before == after


@test
def test_hand_entered_rows_survive_a_rebuild_unchanged():
    """The whole point. A sync must not alter a row a person typed."""
    rows_in = [
        [date(2026, 8, 1), "Acme Parts", "Robot parts", 120.5, "team card", "INV-1", "", "left as typed"],
        [date(2026, 8, 9), "Bolt Supply", "raw materials", 40.0, "R. Coach", "INV-2", "", ""],
    ]
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "rebuilt.xlsx"
        bs.build({}, [], rows_in).save(out)
        reread = bs.read_expenses(openpyxl.load_workbook(out, data_only=True))
    assert len(reread) == 2, reread
    for before, after in zip(rows_in, reread):
        for i, (name, _width) in enumerate(EXPENSE_COLS):
            assert cell_equal(before[i], after[i]), (
                f"{name} changed: {before[i]!r} -> {after[i]!r}"
            )


@test
def test_build_writes_a_categories_sheet_ordered_by_use():
    rows = [
        [date(2026, 8, 1), "Acme Parts", "Tools", 10.0, "", "", "", ""],
        [date(2026, 8, 2), "Bolt Supply", "Robot parts", 20.0, "", "", "", ""],
        [date(2026, 8, 3), "Cog Co", "Robot parts", 30.0, "", "", "", ""],
    ]
    wb = bs.build({}, [], rows)
    assert "Categories" in wb.sheetnames, wb.sheetnames
    ws = wb["Categories"]
    assert [ws.cell(r, 1).value for r in (2, 3)] == ["Robot parts", "Tools"], "most-used first"
    assert [ws.cell(r, 2).value for r in (2, 3)] == [2, 1]


@test
def test_dropdown_covers_the_category_column_and_does_not_forbid_new_values():
    rows = [[date(2026, 8, 1), "Acme Parts", "Tools", 10.0, "", "", "", ""]]
    wb = bs.build({}, [], rows)
    dvs = list(wb["Expenses"].data_validations.dataValidation)
    assert len(dvs) == 1, f"expected one validation, got {len(dvs)}"
    dv = dvs[0]
    assert dv.type == "list"
    assert f"C2:C{bs.CATEGORY_VALIDATION_ROWS}" in str(dv.sqref), str(dv.sqref)
    assert not dv.showErrorMessage, "a new category must not be rejected"
    assert dv.allow_blank, "a blank category must stay allowed"
    assert "Categories" in dv.formula1, dv.formula1
    # showDropDown is inverted in OOXML: true SUPPRESSES the in-cell arrow.
    assert dv.showDropDown in (None, False), (
        f"showDropDown={dv.showDropDown!r} would hide the arrow in Excel"
    )


@test
def test_no_dropdown_when_there_are_no_categories_yet():
    wb = bs.build({}, [], [])
    assert list(wb["Expenses"].data_validations.dataValidation) == [], "empty range is invalid in Excel"


@test
def test_expenses_is_no_longer_blindly_preserved():
    assert "Expenses" not in bs.PRESERVED, (
        "PRESERVED does a values-only copy, which destroys the dropdown and formats"
    )
    assert "Zeffy Map" in bs.PRESERVED and "Sponsors" in bs.PRESERVED


@test
def test_the_script_never_names_the_analysis_workbook():
    """
    An invariant guard, not a behavioural test.

    A behavioural test would need to actually load the analysis workbook to
    prove this script never touches it, and that workbook holds roster data
    that cannot live in a public repo. So this checks the one thing that can
    be checked without one: no *code path* names the file.

    Parsed with ast rather than grepped as raw text, because the filename
    belongs in the documentation — the module docstring's two-file table
    exists specifically to say which file is which — and only a reference
    reachable from running code (a string literal outside a docstring) is
    the defect this guards against. Docstrings are collected first and
    excluded; comments never appear in the AST at all, so they are exempt
    automatically, which is correct for the same reason.
    """
    import ast

    src = (Path(__file__).resolve().parent / "budget_sync.py").read_text()
    tree = ast.parse(src)

    docstring_ids = set()
    for node in ast.walk(tree):
        if isinstance(node, (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)):
            doc_node = node.body[0] if node.body else None
            if (
                isinstance(doc_node, ast.Expr)
                and isinstance(doc_node.value, ast.Constant)
                and isinstance(doc_node.value.value, str)
            ):
                docstring_ids.add(id(doc_node.value))

    offenders = [
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant)
        and isinstance(node.value, str)
        and "FRC Budget" in node.value
        and id(node) not in docstring_ids
    ]
    assert not offenders, [
        f"line {n.lineno}: {n.value!r}" for n in offenders
    ]


def main():
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"  ok    {fn.__name__}")
        except AssertionError as exc:
            failed += 1
            print(f"  FAIL  {fn.__name__}: {exc}")
        except Exception:
            failed += 1
            print(f"  ERROR {fn.__name__}")
            traceback.print_exc()
    print(f"\n{len(TESTS) - failed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
